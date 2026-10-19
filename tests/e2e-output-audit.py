# Output audit: click every download in the app (demo company loaded) and inspect the CONTENT of each file —
# Word/Excel/CSV/ICS/HTML must open, be substantive, and contain no junk (undefined, NaN, [object Object], template
# braces) or unfilled bracket placeholders that Aldune could have filled from the data it has.
import os, sys, re, zipfile, csv, io, json
from lxml import etree
from playwright.sync_api import sync_playwright
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); URL = 'file://' + os.path.join(ROOT, 'dist/paycure-standalone.html')
OUT = os.environ.get('AUDIT_DIR', '/tmp/output-audit'); os.makedirs(OUT, exist_ok=True)
DL = r'^(rpt$|onboard-packet|hb2-issue|tpl2|cl-memo|dl-ics|cal-ics|rd-plan|lit-doc2|notices-docx|dl-|sol-doc|sol-all-docs|bp-csv|bp-xlsx|bp-one|binder|lit-doc|lit-xlsx|pos-memo|sep-agree|xlsx|trail-csv|att-csv|al-csv|dg-docx|board|hb-)'
VIEWS = ['reports', 'documents', 'solutions', 'backpay', 'settlement', 'detective', 'diligence', 'findings', 'litigation', 'pos', 'remediation', 'attest', 'classify', 'calendar', 'readiness', 'notices', 'handbook2', 'agency', 'fixqueue', 'scorecard', 'paga', 'exec', 'method']
JUNK = re.compile(r'\bundefined\b|\bNaN\b|\[object Object\]|\bInfinity\b|\{\{|\}\}|\$-?NaN|\bnull\b(?! and void)|lorem ipsum|TODO', re.I)
PH = re.compile(r'\[(?:[A-Z][A-Za-z ’\'/-]{1,40}|_{2,}|date|name|amount|address)\]')
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
def docx_text(path):
    z = zipfile.ZipFile(path)
    for n in z.namelist():
        if n.endswith('.xml') or n.endswith('.rels'): etree.fromstring(z.read(n))
    root = etree.fromstring(z.read('word/document.xml'))
    return '\n'.join(''.join(t.text or '' for t in p.iter(W + 't')) for p in root.iter(W + 'p'))
def xlsx_text(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=False); out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True): out.append('\t'.join('' if v is None else str(v) for v in row))
    return '\n'.join(out), len(wb.worksheets)
def inspect(path, name):
    ext = name.rsplit('.', 1)[-1].lower(); r = {'file': name, 'ext': ext, 'bytes': os.path.getsize(path), 'issues': []}
    try:
        if ext == 'docx': t = docx_text(path)
        elif ext == 'xlsx': t, n = xlsx_text(path); r['sheets'] = n
        elif ext in ('csv', 'txt', 'md', 'ics', 'html', 'htm', 'json'): t = open(path, encoding='utf8', errors='replace').read()
        elif ext == 'zip':
            z = zipfile.ZipFile(path); t = ' '.join(z.namelist()); r['entries'] = len(z.namelist())
            if not z.namelist(): r['issues'].append('empty zip')
        elif ext == 'pdf': t = open(path, 'rb').read().decode('latin1')[:20000]
        else: t = ''
    except Exception as e:
        r['issues'].append(f'does not open: {str(e)[:80]}'); return r
    if ext == 'html': t = re.sub(r'<script[\s\S]*?</script>|<style[\s\S]*?</style>|<[^>]+>', ' ', t)
    r['chars'] = len(t)
    if ext in ('docx', 'html', 'md', 'txt') and len(t.strip()) < 300: r['issues'].append(f'thin: {len(t.strip())} chars')
    if ext == 'csv' and len([l for l in t.splitlines() if l.strip()]) < 2: r['issues'].append('csv has no data rows')
    if ext == 'ics' and 'BEGIN:VEVENT' not in t: r['issues'].append('calendar has no events')
    j = sorted(set(m.group(0) for m in JUNK.finditer(t)))
    if j: r['issues'].append('junk: ' + ', '.join(j[:6]))
    bl = sorted(set(m.group(1).lower() for m in re.finditer(r'_{6,} \(([^)]{2,60})\)', t)))
    if bl: r['blanks'] = bl
    ph = sorted(set(m.group(0) for m in PH.finditer(t)))
    if ph: r['placeholders'] = ph
    return r
results = []; asked = set()
with sync_playwright() as p:
    b = p.chromium.launch(); pg = b.new_page(viewport={'width': 1400, 'height': 950}, accept_downloads=True); errs = []
    pg.add_init_script('window.__PAYCURE_TEST__=true'); pg.on('pageerror', lambda e: errs.append(f"{pg.evaluate('PCS.view')}: {e}")); pg.on('dialog', lambda d: d.accept())
    pg.goto(URL); pg.wait_for_timeout(800); pg.evaluate("PCX.setView('home')"); pg.wait_for_timeout(300); pg.click('[data-act="demo"] >> nth=0'); pg.wait_for_selector('text=Likely exposure', timeout=60000)
    navs = pg.evaluate("[...document.querySelectorAll('#nav button[data-view]')].map(b => b.dataset.view)")
    seen = set(); k = 0
    for v in list(dict.fromkeys(VIEWS + navs)):
        pg.evaluate(f"PCX.setView('{v}')"); pg.wait_for_timeout(500)
        pg.evaluate("document.querySelectorAll('#views details').forEach(d => d.open = true)")
        n = pg.evaluate(f"[...document.querySelectorAll('#views [data-act]')].filter(x => new RegExp('{DL}').test(x.dataset.act) && x.offsetParent).length")
        for i in range(n):
            pg.evaluate(f"PCX.setView('{v}')"); pg.wait_for_timeout(200); pg.evaluate("document.querySelectorAll('#views details').forEach(d => d.open = true)")
            info = pg.evaluate(f"(()=>{{const el=[...document.querySelectorAll('#views [data-act]')].filter(x => new RegExp('{DL}').test(x.dataset.act) && x.offsetParent)[{i}]; if(!el) return null; return [el.dataset.act, el.dataset.k||el.dataset.kind||el.dataset.r||el.dataset.t||el.dataset.id||'', el.textContent.trim().slice(0,50)]}})()")
            if not info: continue
            key = f"{v}:{info[0]}:{info[1]}:{info[2]}"
            if key in seen: continue
            seen.add(key)
            try:
                with pg.expect_download(timeout=20000) as d:
                    pg.evaluate(f"[...document.querySelectorAll('#views [data-act]')].filter(x => new RegExp('{DL}').test(x.dataset.act) && x.offsetParent)[{i}].click()")
                    pg.wait_for_timeout(150)
                    if pg.locator('#facts-modal').count():  # Aldune asks once for facts the document needs
                        asked.update(pg.evaluate("[...document.querySelectorAll('#facts-modal [data-fact]')].map(i => i.dataset.fact)"))
                        pg.evaluate("""document.querySelectorAll('#facts-modal [data-fact]').forEach(i => { if (i.value) return; const k = i.dataset.fact; i.value = ({ employerAddress: '1450 Harvest Way, Fresno, CA 93721', signerName: 'Dana Whitfield', signerTitle: 'Director of Human Resources', contact: 'Human Resources, (559) 555-0142', caption: 'Garcia v. Pacific Coast Fresh, Inc.', plaintiffs: 'Maria Garcia', caseNumber: '24CV01234', court: 'Superior Court of California, County of Fresno', firm: 'Smith & Lee LLP', attorney: 'Jordan Smith', mediator: 'Hon. Jane Doe (Ret.)', employeeName: 'Luis Ortega', severance: '5000', phoneStipend: '25', commissionPlan: '5% of net collected revenue on all products sold', pieceRate: '$0.45 per tray picked' })[k] || (i.type === 'date' ? '2026-09-15' : 'Test value'); });""")
                        pg.click('#facts-modal [data-fx="save"]')
                k += 1; name = d.value.suggested_filename; fp = os.path.join(OUT, f'{k:03d}-{name}'); d.value.save_as(fp)
                r = inspect(fp, name); r['button'] = key; results.append(r)
            except Exception as e:
                results.append({'button': key, 'file': None, 'issues': [f'no download: {str(e)[:60]}']})
            pg.evaluate("document.querySelectorAll('.drawer.open').forEach(x=>x.classList.remove('open'))")
    b.close()
bad = [r for r in results if r.get('issues')]; ph = [r for r in results if r.get('placeholders')]
print(f'files checked: {len(results)}; with problems: {len(bad)}; with placeholders: {len(ph)}')
for r in bad: print('  PROBLEM', r['button'], '→', r.get('file'), '|', '; '.join(r['issues']))
for r in ph: print('  PLACEHOLDERS', r['button'], '→', r['file'], '|', ' '.join(r['placeholders'][:12]))
import collections
bc = collections.Counter(b for r in results for b in r.get('blanks', []))
print('labeled blanks left for the user (form fields):', dict(bc.most_common(40)))
json.dump(results, open(os.path.join(OUT, 'report.json'), 'w'), indent=1)
print('facts asked for:', sorted(asked))
print('ERRORS', errs[:6])
sys.exit(0 if not bad and not errs else 1)
