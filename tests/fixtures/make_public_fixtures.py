# Synthetic document fixtures for the public repository. Every party, name and term here is invented.
# Builds: Arbitration_Agreement.docx and settlement.pdf (used by document-intel and bulk-import tests).
import os
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

HERE = os.path.dirname(os.path.abspath(__file__))

ARB = [
    ('h', 'Mutual Agreement to Arbitrate Employment-Related Disputes (California)'),
    ('p', 'This Mutual Agreement to Arbitrate ("Agreement") is entered into as of [DATE] between Sample Employer, Inc. (the "Employer") and [EMPLOYEE NAME] (the "Employee").'),
    ('h', '1. Covered Claims'),
    ('p', 'Except as stated in Section 3, the Employer and the Employee agree to resolve by final and binding arbitration any dispute arising out of the Employee\'s employment, including claims for wages, overtime, meal and rest period premiums, wage statement penalties, waiting-time penalties, discrimination, harassment, and retaliation.'),
    ('h', '2. Arbitration Procedure'),
    ('p', 'Arbitration will be conducted by a single neutral arbitrator under the employment arbitration rules of a recognized provider. The arbitrator will apply the substantive law of California, permit adequate discovery, and issue a written decision. The Employer will pay all costs unique to arbitration, including the arbitrator\'s fees.'),
    ('h', '3. Excluded Claims'),
    ('p', 'This Agreement does not cover claims for workers\' compensation or unemployment insurance benefits, claims that may not be arbitrated as a matter of law, or claims of sexual harassment or sexual assault that the Employee elects to bring in court under federal law.'),
    ('h', '4. Class, Collective, and Representative Actions'),
    ('p', 'Claims will be arbitrated on an individual basis only. The Employee\'s individual claims under the Private Attorneys General Act (PAGA), Labor Code section 2699, will be arbitrated individually. Representative PAGA claims brought on behalf of other employees are not waived; they will be stayed in the Superior Court pending the outcome of the individual arbitration.'),
    ('h', '5. Voluntary Agreement'),
    ('p', 'The Employee has had the opportunity to review this Agreement and consult counsel. This Agreement survives the end of employment and may be modified only in a writing signed by both parties.'),
]

def build_arb():
    d = Document(); d.core_properties.author = 'Aldune test data'; d.core_properties.last_modified_by = 'Aldune test data'
    for kind, text in ARB:
        (d.add_heading(text, level=1) if kind == 'h' else d.add_paragraph(text))
    d.save(os.path.join(HERE, 'Arbitration_Agreement.docx'))

SETTLEMENT = [
    'SETTLEMENT AGREEMENT AND GENERAL RELEASE',
    '',
    'This Settlement Agreement and General Release ("Agreement") is entered into between Sample Hospitality, LLC',
    '("Company") and Jordan Example ("Employee"). The parties are fictional and this document is test data.',
    '',
    '1. Payment. Company will pay Employee the gross sum of $12,500, less applicable withholdings,',
    '   in full satisfaction of all claims for unpaid wages, overtime, and meal and rest period premiums.',
    '',
    '2. Release. Employee releases Company from all claims arising from Employee\'s employment, known or',
    '   unknown, and waives the protections of Civil Code section 1542.',
    '',
    '3. Confidentiality. Employee agrees to keep the terms and amount of this Agreement confidential and',
    '   not to discuss them with any third party.',
    '',
    '4. No Admission. This Agreement is not an admission of liability by Company.',
    '',
    '5. Entire Agreement. This Agreement is the entire agreement between the parties on its subject.',
    '',
    'Signed: ______________________ (Employee)      ______________________ (Company)',
]

def build_settlement():
    c = canvas.Canvas(os.path.join(HERE, 'settlement.pdf'), pagesize=letter); c.setAuthor('Aldune test data'); c.setTitle('Sample settlement agreement')
    y = 740
    for line in SETTLEMENT:
        c.drawString(56, y, line); y -= 16
    c.showPage(); c.drawString(56, 740, 'Exhibit A: payment schedule (test data).'); c.save()

if __name__ == '__main__':
    build_arb(); build_settlement(); print('built Arbitration_Agreement.docx, settlement.pdf')
